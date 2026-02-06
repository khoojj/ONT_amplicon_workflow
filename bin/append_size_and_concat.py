#!/usr/bin/env python3
import argparse
import re
from pathlib import Path
from openpyxl import load_workbook


def cluster_num(consensus_id: str) -> int:
    m = re.search(r'c([0-9]+)$', consensus_id)
    return int(m.group(1)) if m else 10**9


def read_excel_map(xlsx_path: str):
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    header_row = next(ws.iter_rows(min_row=1, max_row=1))
    hdr = [str(c.value).strip() if c.value is not None else "" for c in header_row]
    col = {name: i for i, name in enumerate(hdr)}

    need = ["consensus_id", "reads_in_cluster", "draft_id"]
    for n in need:
        if n not in col:
            raise SystemExit(f"Missing column '{n}' in {xlsx_path}. Have: {hdr}")

    m_reads = {}
    m_uuid = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        cid = row[col["consensus_id"]]
        if cid is None:
            continue
        cid = str(cid).strip()

        reads = row[col["reads_in_cluster"]]
        draft = row[col["draft_id"]]
        uuid = str(draft).split()[0] if draft is not None else "NA"

        m_reads[cid] = int(reads) if reads is not None else 0
        m_uuid[cid] = uuid

    return m_reads, m_uuid


def infer_ids_from_fastas(fastas):
    ids = []
    for fa in fastas:
        stem = Path(fa).name
        # remove one extension (e.g. ".fasta" or ".fa")
        cid = Path(stem).stem
        ids.append(cid)
    return ids


def main():
    ap = argparse.ArgumentParser(
        description="Append reads_in_cluster as ;size= to FASTA headers and concatenate per-barcode consensus FASTAs."
    )
    ap.add_argument("--xlsx", required=True, help="Excel file produced from nanoclust_out.txt")
    ap.add_argument("--out", required=True, help="Output FASTA (e.g., b06_consensus.fasta)")
    ap.add_argument("--fastas", nargs="+", required=True, help="Input consensus fasta files (e.g., b06c0.fasta ...)")
    ap.add_argument(
        "--ids",
        nargs="+",
        default=None,
        help="Consensus IDs corresponding to --fastas (optional). If omitted, inferred from FASTA filenames."
    )
    args = ap.parse_args()

    fastas = args.fastas
    ids = args.ids if args.ids is not None else infer_ids_from_fastas(fastas)

    if len(fastas) != len(ids):
        raise SystemExit(f"--fastas ({len(fastas)}) and --ids ({len(ids)}) must have the same length")

    m_reads, m_uuid = read_excel_map(args.xlsx)

    pairs = list(zip(ids, fastas))
    pairs.sort(key=lambda x: cluster_num(x[0]))

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w") as oh:
        for cid, fa in pairs:
            if cid not in m_reads:
                raise SystemExit(
                    f"Consensus id '{cid}' not found in Excel column 'consensus_id'. "
                    f"(FASTA file was '{fa}')"
                )

            reads = m_reads[cid]
            uuid = m_uuid.get(cid, "NA")

            oh.write(f">{cid}_{uuid};size={reads};\n")

            with open(fa, "r") as fh:
                for line in fh:
                    if line.startswith(">"):
                        continue
                    oh.write(line)

    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()
